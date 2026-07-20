"""
Tạo file Excel mẫu cho một chiến dịch mới.

Chạy:
    python tao_file_mau.py tuyen_dung.xlsx TuyenDung
    python tao_file_mau.py seeding_web.xlsx Seeding

Tham số: <tên file> <tên sheet>. File tạo ra có sẵn dòng tiêu đề và 2 dòng ví
dụ để bạn xóa đi rồi điền dữ liệu thật. Cấu trúc giống hệt phong_tro.xlsx:
  A: STT | B: Nội dung bài | C: Phân loại | D: Tên thư mục ảnh
"""
import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

HEADER = ["STT", "Noi dung bai (Caption Facebook)", "Phan loai", "Folder anh"]
WIDTHS = {"A": 8, "B": 60, "C": 25, "D": 20}

VI_DU = [
    (1, "Nội dung bài đăng số 1.\nXuống dòng thoải mái, đăng lên Facebook giữ nguyên.",
     "phan loai a", "thu_muc_anh_1"),
    (2, "Nội dung bài đăng số 2.\nĐể trống cột D nếu bài này không cần ảnh.",
     "phan loai b", ""),
]


def tao_file(ten_file, ten_sheet):
    path = ten_file if os.path.isabs(ten_file) else os.path.join(BASE_DIR, ten_file)
    if os.path.exists(path):
        print(f"⚠ '{path}' đã tồn tại — không ghi đè. Đổi tên khác hoặc xóa file cũ.")
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ten_sheet

    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in VI_DU:
        ws.append(row)

    # Nội dung nhiều dòng nên bật wrap để nhìn được trong Excel
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width

    wb.save(path)
    print(f"✓ Đã tạo {path}  (sheet '{ten_sheet}')")
    return path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    tao_file(sys.argv[1], sys.argv[2])
